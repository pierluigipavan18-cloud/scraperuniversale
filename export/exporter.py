"""Export module - CSV, JSON, Excel output."""

from __future__ import annotations

import csv
import json
import logging
import os
from datetime import datetime
from pathlib import Path

from scraper.models import Company

logger = logging.getLogger("export.exporter")

FIELDS = [
    "name", "city", "country", "address", "email", "phone",
    "fax", "website", "sector", "description", "employees",
    "vat_id", "source", "source_url", "emails_extra",
]


def _ensure_dir(output_dir: str) -> Path:
    path = Path(output_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _filename(output_dir: str, prefix: str, ext: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = _ensure_dir(output_dir)
    return str(path / f"{prefix}_{ts}.{ext}")


def export_csv(
    companies: list[Company],
    output_dir: str = "./output",
    prefix: str = "companies",
) -> str:
    """Export companies to CSV file."""
    filepath = _filename(output_dir, prefix, "csv")

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        for company in companies:
            writer.writerow(company.to_dict())

    logger.info(f"Exported {len(companies)} companies to {filepath}")
    return filepath


def export_json(
    companies: list[Company],
    output_dir: str = "./output",
    prefix: str = "companies",
) -> str:
    """Export companies to JSON file."""
    filepath = _filename(output_dir, prefix, "json")

    data = [company.to_dict() for company in companies]
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    logger.info(f"Exported {len(companies)} companies to {filepath}")
    return filepath


def export_excel(
    companies: list[Company],
    output_dir: str = "./output",
    prefix: str = "companies",
) -> str:
    """Export companies to Excel file."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        logger.error("openpyxl not installed. Run: pip install openpyxl")
        raise

    filepath = _filename(output_dir, prefix, "xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"

    # Header style
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

    # Write headers
    for col, field in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=col, value=field.replace("_", " ").title())
        cell.font = header_font
        cell.fill = header_fill

    # Write data
    for row, company in enumerate(companies, 2):
        data = company.to_dict()
        for col, field in enumerate(FIELDS, 1):
            ws.cell(row=row, column=col, value=data.get(field, ""))

    # Auto-width columns
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        adjusted = min(max_length + 2, 50)
        ws.column_dimensions[col[0].column_letter].width = adjusted

    wb.save(filepath)
    logger.info(f"Exported {len(companies)} companies to {filepath}")
    return filepath


def export(
    companies: list[Company],
    fmt: str = "csv",
    output_dir: str = "./output",
    prefix: str = "companies",
) -> str:
    """Export companies in the specified format."""
    exporters = {
        "csv": export_csv,
        "json": export_json,
        "excel": export_excel,
        "xlsx": export_excel,
    }

    exporter = exporters.get(fmt.lower())
    if not exporter:
        raise ValueError(f"Unknown format: {fmt}. Use: {', '.join(exporters.keys())}")

    return exporter(companies, output_dir, prefix)
